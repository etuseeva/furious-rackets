#!/usr/bin/env python3
"""Собирает content-template.xlsx из site/js/default-content.js.

Запуск из корня проекта:
    python3 -m pip install openpyxl
    python3 tools/build_template.py
"""

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "site" / "js" / "default-content.js"
TARGET = ROOT / "content-template.xlsx"

LAST_FORMATTED_ROW = 100

COLUMNS = {
    "Настройки": ["ключ", "значение", "пояснение"],
    "Тренировки": ["название", "описание", "показывать"],
    "Расписание": ["день", "время", "площадка", "тип", "уровни", "цена", "показывать"],
    "Площадки": ["название", "полное название", "адрес", "ссылка на карту", "показывать"],
    "Акции": ["заголовок", "текст", "картинка", "текст кнопки", "ссылка кнопки", "до", "показывать"],
    "Турниры": ["название", "даты", "место", "текст", "афиша", "ссылка", "до", "показывать"],
    "Тренер": ["имя", "фото", "о себе", "telegram", "показывать"],
    "Партнёры": ["название", "описание", "логотип", "ссылка", "показывать"],
    "FAQ": ["вопрос", "ответ", "показывать"],
    "Соцсети": ["название", "ссылка", "показывать"],
}

WIDTHS = {
    "текст": 70,
    "ответ": 70,
    "о себе": 50,
    "вопрос": 46,
    "значение": 50,
    "пояснение": 46,
    "полное название": 46,
    "адрес": 42,
    "заголовок": 32,
    "название": 34,
    "описание": 70,
    "ссылка": 28,
    "ссылка кнопки": 26,
    "ссылка на карту": 26,
    "текст кнопки": 24,
    "картинка": 24,
    "афиша": 24,
    "логотип": 24,
    "telegram": 24,
    "фото": 24,
    "даты": 28,
    "место": 20,
    "уровни": 16,
    "время": 16,
    "день": 16,
    "имя": 22,
    "ключ": 22,
    "цена": 10,
    "тип": 14,
    "до": 14,
    "показывать": 13,
    "площадка": 18,
}

HEADER_FILL = PatternFill("solid", fgColor="B6F024")
HEADER_FONT = Font(bold=True, color="000000")

INSTRUCTIONS = [
    ("Как заполнять таблицу", True),
    ("", False),
    ("Эта таблица — источник содержимого сайта «Бешеные ракетки».", False),
    ("Сайт перечитывает её при каждом открытии страницы. Изменения видны примерно через минуту.", False),
    ("", False),
    ("Главные правила", True),
    ("1. Не переименовывайте листы и не меняйте текст в первой строке (заголовки столбцов).", False),
    ("   Сайт ищет листы и столбцы по именам. Переименовали — блок вернётся к стартовому содержимому.", False),
    ("2. Столбцы можно оставлять пустыми, строки — добавлять и удалять.", False),
    ("3. Полностью пустые строки сайт пропускает.", False),
    ("", False),
    ("Столбец «показывать»", True),
    ("Пусто или «да» — строка видна на сайте.", False),
    ("«нет», «no», «0», «false» — строка скрыта.", False),
    ("Если на листе не осталось видимых строк, весь блок и пункт меню исчезают с сайта.", False),
    ("", False),
    ("Столбец «до»", True),
    ("Дата в формате 31.12.2026 или 2026-12-31.", False),
    ("Строка пропадает с сайта на следующий день после этой даты.", False),
    ("Пусто — показывается бессрочно.", False),
    ("", False),
    ("Расписание и площадки", True),
    ("В столбце «площадка» пишите короткое название ровно так, как оно записано", False),
    ("в столбце «название» на листе «Площадки» — сайт связывает их по этому полю", False),
    ("и подставляет в расписание полное название.", False),
    ("Столбец «уровни»: перечисляйте через запятую — B, C+, E. Пусто — на сайте будет «все уровни».", False),
    ("Столбец «тип»: «рабочая» или «игровая».", False),
    ("Столбец «цена»: только число, знак ₽ сайт добавит сам.", False),
    ("«ссылка на карту» можно не заполнять — сайт сам построит ссылку на Яндекс Карты по адресу.", False),
    ("", False),
    ("Картинки", True),
    ("Картинки лежат в папке assets репозитория. В таблице пишите путь к файлу:", False),
    ("assets/promo-pool.jpg, assets/tournament.jpg и так далее.", False),
    ("Можно указать и полный адрес картинки в интернете (https://…).", False),
    ("", False),
    ("Ссылки внутри текста", True),
    ("В любом тексте (ответы FAQ, описания тренировок, акции, турниры) можно поставить:", False),
    ("{телефон} — подставится номер из настроек, по нему можно позвонить в один клик;", False),
    ("{тренер} — ник тренера, ссылка на его Telegram;", False),
    ("{группа} — ссылка на группу клуба в Telegram.", False),
    ("Сами адреса писать не нужно — сайт подставит их из «Настроек».", False),
    ("", False),
    ("Ссылки в отдельных столбцах", True),
    ("Принимаются только адреса, начинающиеся на https://, http://, tel:, mailto:,", False),
    ("а также пути внутри сайта (assets/…). Остальное сайт проигнорирует.", False),
    ("", False),
    ("Лист «Настройки»", True),
    ("Столбец «ключ» не трогайте — меняйте только «значение».", False),
    ("phone — телефон главной кнопки, phone_name — подпись под ним (кто возьмёт трубку).", False),
    ("group_telegram — группа, куда ведёт кнопка «Группа в Telegram».", False),
    ("Остальные ссылки — ВКонтакте, Instagram, MAX, чат — живут на листе «Соцсети».", False),
    ("", False),
    ("Формат ячеек", True),
    ("Во всех листах ячейки заранее переведены в текстовый формат, чтобы Google Таблицы", False),
    ("не превращали цены и даты в числа. Новые строки копируйте из уже заполненных.", False),
]


def load_content():
    text = SOURCE.read_text(encoding="utf-8")
    match = re.search(r"window\.DEFAULT_CONTENT\s*=\s*(\{.*\})\s*;\s*$", text, re.S)
    if not match:
        raise SystemExit(f"Не удалось разобрать {SOURCE}: ожидается window.DEFAULT_CONTENT = {{...}};")
    return json.loads(match.group(1))


def write_instructions(sheet):
    sheet.column_dimensions["A"].width = 110
    for index, (line, is_title) in enumerate(INSTRUCTIONS, start=1):
        cell = sheet.cell(row=index, column=1, value=line)
        cell.alignment = Alignment(vertical="center")
        if is_title:
            cell.font = Font(bold=True, size=13 if index == 1 else 11)
            if index > 1:
                cell.fill = HEADER_FILL


def write_sheet(sheet, columns, rows):
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS.get(name, 22)

    for row_index, row in enumerate(rows, start=2):
        for column_index, name in enumerate(columns, start=1):
            sheet.cell(row=row_index, column=column_index, value=row.get(name, ""))

    for row_index in range(2, LAST_FORMATTED_ROW + 1):
        for column_index in range(1, len(columns) + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"


def main():
    content = load_content()

    workbook = Workbook()
    write_instructions(workbook.active)
    workbook.active.title = "Как заполнять"

    for sheet_name, columns in COLUMNS.items():
        rows = content.get(sheet_name, [])
        write_sheet(workbook.create_sheet(sheet_name), columns, rows)

    workbook.save(TARGET)
    print(f"Готово: {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
